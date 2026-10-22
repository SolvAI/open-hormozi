"""Markdown → DOCX processor. Minimal, crash-proof."""

from __future__ import annotations

import re
import logging
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from engine.processors.base import FormatProcessor

logger = logging.getLogger(__name__)

_HEADING_PATTERN = re.compile(r"^(#{1,4})\s+(.+)$")
_BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_PATTERN = re.compile(r"\*(.+?)\*")
_BULLET_PATTERN = re.compile(r"^[-*]\s+(.+)$")
_NUMBERED_PATTERN = re.compile(r"^\d+\.\s+(.+)$")
_TABLE_SEP_PATTERN = re.compile(r"^\|[\s\-:|]+\|$")
_TABLE_ROW_PATTERN = re.compile(r"^\|(.+)\|$")
_CODE_BLOCK_FENCE = re.compile(r"^```")


class DocProcessor(FormatProcessor):

    def process(
        self,
        content: Any,
        doc_id: str,
        output_dir: Path,
        template_path: str | None = None,
        metadata: dict | None = None,
    ) -> Path:
        self._ensure_dir(output_dir)
        md_text = content if isinstance(content, str) else str(content)

        doc = Document()

        doc_name = (metadata or {}).get("doc_name", doc_id)
        doc.add_heading(doc_name, level=0)

        self._convert_markdown(doc, md_text)

        filename = f"{doc_id}_{self._slugify(doc_name)}.docx"
        output_path = output_dir / filename
        doc.save(str(output_path))
        logger.info(f"DocProcessor: wrote {output_path}")
        return output_path

    def _convert_markdown(self, doc: Document, md_text: str) -> None:
        lines = md_text.split("\n")
        i = 0
        in_code_block = False
        code_buffer: list[str] = []

        while i < len(lines):
            line = lines[i]

            # Code fences
            if _CODE_BLOCK_FENCE.match(line.strip()):
                if in_code_block:
                    if code_buffer:
                        p = doc.add_paragraph()
                        run = p.add_run("\n".join(code_buffer))
                        run.font.name = "Consolas"
                        run.font.size = Pt(9)
                    code_buffer = []
                    in_code_block = False
                else:
                    in_code_block = True
                i += 1
                continue

            if in_code_block:
                code_buffer.append(line)
                i += 1
                continue

            # Empty lines
            if not line.strip():
                i += 1
                continue

            # Headings
            heading_match = _HEADING_PATTERN.match(line)
            if heading_match:
                level = min(len(heading_match.group(1)), 4)
                doc.add_heading(heading_match.group(2), level=level)
                i += 1
                continue

            # Tables
            if _TABLE_ROW_PATTERN.match(line.strip()):
                table_lines = []
                while i < len(lines) and _TABLE_ROW_PATTERN.match(lines[i].strip()):
                    if not _TABLE_SEP_PATTERN.match(lines[i].strip()):
                        table_lines.append(lines[i].strip())
                    i += 1
                if table_lines:
                    self._add_table(doc, table_lines)
                continue

            # Bullets
            bullet_match = _BULLET_PATTERN.match(line)
            if bullet_match:
                self._add_rich_paragraph(doc, bullet_match.group(1), style="List Bullet")
                i += 1
                continue

            # Numbered
            num_match = _NUMBERED_PATTERN.match(line)
            if num_match:
                self._add_rich_paragraph(doc, num_match.group(1), style="List Number")
                i += 1
                continue

            # Regular paragraph
            self._add_rich_paragraph(doc, line)
            i += 1

    def _add_rich_paragraph(self, doc: Document, text: str, style: str | None = None) -> None:
        """Add paragraph with bold/italic support. Falls back to no style if style missing."""
        try:
            p = doc.add_paragraph(style=style)
        except KeyError:
            p = doc.add_paragraph()

        parts = _BOLD_PATTERN.split(text)
        for idx, part in enumerate(parts):
            if not part:
                continue
            if idx % 2 == 1:
                run = p.add_run(part)
                run.bold = True
            else:
                italic_parts = _ITALIC_PATTERN.split(part)
                for j, ipart in enumerate(italic_parts):
                    if not ipart:
                        continue
                    run = p.add_run(ipart)
                    if j % 2 == 1:
                        run.italic = True

    def _add_table(self, doc: Document, rows: list[str]) -> None:
        parsed = []
        for row in rows:
            cells = [c.strip() for c in row.strip("|").split("|")]
            parsed.append(cells)
        if not parsed:
            return

        num_cols = len(parsed[0])
        table = doc.add_table(rows=len(parsed), cols=num_cols)

        try:
            table.style = "Table Grid"
        except KeyError:
            pass

        for i, row_data in enumerate(parsed):
            for j, cell_text in enumerate(row_data):
                if j < num_cols:
                    table.cell(i, j).text = cell_text

    @staticmethod
    def _slugify(text: str) -> str:
        slug = text.lower().replace(" ", "_").replace("&", "and")
        return re.sub(r"[^a-z0-9_]", "", slug)